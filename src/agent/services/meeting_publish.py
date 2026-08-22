"""Publish meeting plans to SharePoint (Markdown + formatted DOCX)."""

from __future__ import annotations

import re
from datetime import datetime, timezone
from typing import Any

from agent.config import settings
from agent.core.logging import get_logger
from agent.services import graph_docs

logger = get_logger(__name__)


def _slug(text: str, *, max_len: int = 48) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "-", (text or "").lower()).strip("-")
    return (s[:max_len] or "plan").strip("-")


def plan_filenames(plan: dict[str, Any], *, plan_kind: str = "implementation") -> tuple[str, str]:
    date = datetime.now(timezone.utc).strftime("%Y-%m-%d")
    slug = _slug(plan.get("title") or plan.get("topic_label") or "plan")
    suffix = "action-plan" if plan_kind == "action" else "plan"
    base = f"{date}-{slug}-{suffix}"
    return f"{base}.md", f"{base}.docx"


def render_docx_bytes(
    plan: dict[str, Any],
    markdown: str,
    *,
    client_name: str = "",
    plan_kind: str = "implementation",
) -> bytes | None:
    """Render a formatted DOCX when python-docx is available."""
    try:
        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        logger.info("python-docx not installed — MD-only publish")
        return None

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)

    title = plan.get("title") or ("Meeting Action Plan" if plan_kind == "action" else "Implementation Plan")
    subtitle_parts = []
    if client_name:
        subtitle_parts.append(f"Client: {client_name}")
    subtitle_parts.append(datetime.now(timezone.utc).strftime("%d %B %Y"))
    if plan_kind == "action":
        subtitle_parts.append("Action Plan")
    else:
        subtitle_parts.append("Implementation Plan")

    h = doc.add_heading(title, level=0)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    sub = doc.add_paragraph(" · ".join(subtitle_parts))
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    if sub.runs:
        sub.runs[0].font.size = Pt(11)
        sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_paragraph("")

    def _add_table_from_markdown_table(lines: list[str]) -> None:
        if len(lines) < 2:
            return
        headers = [c.strip() for c in lines[0].strip("|").split("|")]
        rows = []
        for line in lines[2:]:
            if line.strip():
                rows.append([c.strip() for c in line.strip("|").split("|")])
        if not headers:
            return
        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        for i, label in enumerate(headers):
            hdr_cells[i].text = label
            for p in hdr_cells[i].paragraphs:
                for r in p.runs:
                    r.bold = True
        for ri, row in enumerate(rows):
            for ci, val in enumerate(row[: len(headers)]):
                table.rows[ri + 1].cells[ci].text = val
        doc.add_paragraph("")

    table_buf: list[str] = []
    for line in markdown.splitlines():
        text = line.rstrip()
        if text.startswith("|"):
            table_buf.append(text)
            continue
        if table_buf:
            _add_table_from_markdown_table(table_buf)
            table_buf = []
        if not text or text.startswith("# "):
            if text.startswith("# "):
                doc.add_heading(text[2:].strip(), level=1)
            continue
        if text.startswith("## "):
            doc.add_heading(text[3:].strip(), level=2)
        elif text.startswith("### "):
            doc.add_heading(text[4:].strip(), level=3)
        elif text.startswith("- "):
            doc.add_paragraph(text[2:], style="List Bullet")
        else:
            p = doc.add_paragraph(text)
            if p.runs and text == (plan.get("executive_summary") or ""):
                p.runs[0].italic = True

    if table_buf:
        _add_table_from_markdown_table(table_buf)

    import io

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def publish_plan(
    plan: dict[str, Any],
    markdown: str,
    *,
    folder: str | None = None,
    plan_kind: str = "implementation",
    client_name: str = "",
) -> dict[str, Any]:
    target_folder = (folder or settings.meeting_plans_folder or "MeetingPlans").strip().strip("/")
    md_name, docx_name = plan_filenames(plan, plan_kind=plan_kind)
    formats = (settings.meeting_plan_default_format or "md,docx").lower()

    md_url = ""
    docx_url = ""
    if "md" in formats:
        entry = await graph_docs.upload_site_drive_item(
            folder=target_folder,
            filename=md_name,
            content=markdown.encode("utf-8"),
            content_type="text/markdown; charset=utf-8",
        )
        md_url = (entry or {}).get("web_url") or ""

    docx_bytes = render_docx_bytes(
        plan,
        markdown,
        client_name=client_name,
        plan_kind=plan_kind,
    ) if "docx" in formats else None
    if docx_bytes:
        entry = await graph_docs.upload_site_drive_item(
            folder=target_folder,
            filename=docx_name,
            content=docx_bytes,
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        docx_url = (entry or {}).get("web_url") or ""

    return {
        "title": plan.get("title") or "Meeting Plan",
        "md_url": md_url,
        "docx_url": docx_url,
        "md_filename": md_name,
        "docx_filename": docx_name,
        "folder": target_folder,
        "plan_kind": plan_kind,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
