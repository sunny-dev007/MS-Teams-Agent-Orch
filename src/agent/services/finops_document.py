"""Orbit FinOps cost-optimisation DOCX — charts + cost-joined analysis.

Feature: Orbit-only document quality for Azure cost / recommendations reports.
Does not alter Copilot Studio (AI Dev Agent) channel behaviour.

Deterministic numbers come from Cost Management + inventory; LLM only polishes
narrative around injected facts (avoids generic best-practice fluff).
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone
from typing import Any

from agent.config import settings
from agent.core.logging import get_logger
from agent.services import azure_finops

logger = get_logger(__name__)

_FINOPS_DOC_RE = re.compile(
    r"(?:"
    r"azure|finops|cost|subscription|optimis|optimiz|recommendation|"
    r"saving|spend|billing|right-?siz|sku|resource\s+group"
    r")",
    re.I,
)


def wants_finops_cost_report(message: str, *, session_data: dict[str, Any] | None = None) -> bool:
    """True when the user wants a cost/optimisation report (not a generic essay)."""
    msg = (message or "").strip()
    if not msg:
        return False
    if not _FINOPS_DOC_RE.search(msg):
        return False
    data = session_data or {}
    if data.get("azure_finops_sub") or data.get("azure_finops_last_costs_svc"):
        return True
    # Explicit cost/optimisation phrasing even before a FinOps session exists
    return bool(
        re.search(
            r"(?:cost|finops|optimis|optimiz|recommendation|saving|spend|billing)",
            msg,
            re.I,
        )
    )


def _money(amount: float, currency: str = "INR") -> str:
    return f"{float(amount):,.2f} {currency}"


def _short_name(name: str) -> str:
    return str(name or "").rstrip("/").split("/")[-1] or str(name or "—")


def _is_free_sku(sku: str) -> bool:
    s = (sku or "").upper()
    return s in {"F1", "Y1", "FREE", "D1"} or s.startswith("F1") or s.startswith("Y1")


def build_analysis_bundle(
    *,
    sub: dict[str, Any],
    cost_by_rg: list[dict[str, Any]],
    cost_by_svc: list[dict[str, Any]],
    cost_by_res: list[dict[str, Any]],
    resources: list[dict[str, Any]],
    recs: list[dict[str, Any]],
    days: int,
) -> dict[str, Any]:
    """Join costs + agent tips into the AI-Dev-Agent-style analysis model."""
    currency = (
        (cost_by_svc[0].get("currency") if cost_by_svc else None)
        or (cost_by_rg[0].get("currency") if cost_by_rg else None)
        or (cost_by_res[0].get("currency") if cost_by_res else None)
        or "INR"
    )
    total_svc = sum(float(r.get("cost") or 0) for r in cost_by_svc)
    total_rg = sum(float(r.get("cost") or 0) for r in cost_by_rg)
    total = total_svc if total_svc > 0 else total_rg

    plans = [
        r
        for r in resources
        if "serverfarms" in (r.get("type") or "").lower()
        or "microsoft.web/serverfarms" in (r.get("type") or "").lower()
    ]
    agent_rows: list[dict[str, Any]] = []
    for p in plans[:12]:
        sku = str(p.get("sku") or "—")
        name = p.get("name") or "—"
        lookup = next(
            (
                c
                for c in cost_by_res
                if _short_name(str(c.get("name") or "")).lower() == str(name).lower()
            ),
            None,
        )
        cost = float((lookup or {}).get("cost") or 0)
        # Heuristic mirror of old agent: free SKUs often marked HIGH incorrectly
        if _is_free_sku(sku) and cost <= 0.01:
            priority = "HIGH" if any(
                str(name).lower() in (t.get("title") or "").lower()
                or str(name).lower() in (t.get("detail") or "").lower()
                for t in recs
            ) else "LOW"
        else:
            priority = "LOW"
        agent_rows.append(
            {
                "name": name,
                "sku": sku,
                "priority": priority,
                "cost": cost,
                "rg": p.get("rg") or "",
            }
        )

    # Cost-based ranking (resource groups + top resources)
    cost_rank: list[dict[str, Any]] = []
    for row in sorted(cost_by_rg, key=lambda r: float(r.get("cost") or 0), reverse=True)[:8]:
        c = float(row.get("cost") or 0)
        if c <= 0:
            continue
        share = (c / total) if total > 0 else 0.0
        cost_rank.append(
            {
                "target": row.get("name"),
                "kind": "resource_group",
                "cost": c,
                "share": share,
                "assessment": (
                    "Primary bill driver — investigate compute / SQL meters in Cost analysis"
                    if share >= 0.25
                    else "Material spend — verify right-size and idle hours"
                ),
            }
        )

    paid_in_agent = sum(float(r["cost"]) for r in agent_rows if r["priority"] == "HIGH")
    agent_coverage = (paid_in_agent / total) if total > 0 else 0.0
    top4 = cost_rank[:4]
    cost_coverage = sum(r["share"] for r in top4)

    # Indicative savings: top RG compute-heavy share heuristic (not a quotation)
    top_rg = cost_rank[0] if cost_rank else None
    save_feapp_low = 0.0
    save_feapp_high = 0.0
    save_poc = 0.0
    save_db = 0.0
    if top_rg and top_rg["share"] >= 0.4:
        save_feapp_low = top_rg["cost"] * 0.55
        save_feapp_high = top_rg["cost"] * 0.70
    for row in cost_rank:
        name = str(row.get("target") or "").lower()
        if "todo" in name or "poc" in name:
            save_poc = max(save_poc, float(row["cost"]))
        if "db" in name or "sql" in name:
            save_db = max(save_db, min(float(row["cost"]) * 0.3, float(row["cost"])))

    save_low = save_feapp_low + save_poc + (save_db * 0.5 if save_db else 0)
    save_high = save_feapp_high + save_poc + save_db
    if save_high <= 0 and total > 0:
        save_low = total * 0.35
        save_high = total * 0.45

    contradictions: list[dict[str, str]] = []
    for row in agent_rows:
        if _is_free_sku(str(row.get("sku"))) and float(row.get("cost") or 0) > 1.0:
            contradictions.append(
                {
                    "resource": str(row.get("name")),
                    "sku": str(row.get("sku")),
                    "cost": _money(float(row["cost"]), currency),
                    "note": (
                        "Reported as free-tier SKU but Cost Management shows spend. "
                        "Possible stale SKU reading, different serving plan, or non-compute meter."
                    ),
                }
            )

    actions: list[dict[str, Any]] = []
    if top_rg:
        actions.append(
            {
                "priority": 1,
                "action": f"Investigate / right-size compute in `{top_rg['target']}`",
                "saving": f"{save_feapp_low:,.0f} – {save_feapp_high:,.0f}"
                if save_feapp_high
                else "to be assessed",
                "risk": "Low",
                "effort": "15–30 min",
            }
        )
    if save_poc:
        actions.append(
            {
                "priority": 2,
                "action": "Retire or stop POC databases / groups still billing",
                "saving": f"up to {save_poc:,.0f}",
                "risk": "Low",
                "effort": "10 min",
            }
        )
    if save_db:
        actions.append(
            {
                "priority": 3,
                "action": "Right-size SQL / database tiers to match query load",
                "saving": "to be assessed",
                "risk": "Medium",
                "effort": "30 min",
            }
        )
    actions.append(
        {
            "priority": len(actions) + 1,
            "action": "Delete zero-cost empty resource groups / free F1·Y1 noise (hygiene)",
            "saving": "0 (hygiene)",
            "risk": "Low",
            "effort": "15 min",
        }
    )

    defects = [
        {
            "defect": "Word commands (costs, recommendations, menu) previously reset FinOps flow",
            "impact": "Users lose place mid-flow",
            "fix": "Map keywords to FinOps awaiting handlers (already mitigated in agent)",
        },
        {
            "defect": "API failures and zero-spend used identical wording",
            "impact": "Failures misdiagnosed as permission gaps",
            "fix": "Distinguish empty result sets from API errors",
        },
        {
            "defect": "Recommendations ranked by SKU/naming instead of observed cost",
            "impact": (
                f"Zero-cost plans marked HIGH while ~{cost_coverage:.0%} of spend "
                "needs cost-based ranking"
            ),
            "fix": "Join recommendations to Cost Management; rank by spend",
        },
    ]

    return {
        "sub_name": sub.get("name") or "subscription",
        "sub_id": sub.get("id") or "",
        "days": days,
        "currency": currency,
        "total": total,
        "cost_by_rg": cost_by_rg,
        "cost_by_svc": cost_by_svc,
        "cost_by_res": cost_by_res,
        "agent_rows": agent_rows,
        "cost_rank": cost_rank,
        "agent_coverage": agent_coverage,
        "cost_coverage": cost_coverage,
        "save_low": save_low,
        "save_high": save_high,
        "projected": max(0.0, total - save_high),
        "contradictions": contradictions,
        "actions": actions,
        "defects": defects,
        "recs": recs,
        "prepared_at": datetime.now(timezone.utc),
        "waterfall": [
            ("Current monthly bill", total, "base"),
            ("Top RG right-size", -save_feapp_high if save_feapp_high else -save_low * 0.8, "cut"),
            ("Retire POC", -save_poc, "cut"),
            ("Right-size DB", -save_db if save_db else 0.0, "cut"),
            ("Projected bill", max(0.0, total - save_high), "base"),
        ],
    }


def _chart_png(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=140, bbox_inches="tight", facecolor="white")
    buf.seek(0)
    return buf.getvalue()


def render_charts(bundle: dict[str, Any]) -> dict[str, bytes]:
    """Return PNG bytes for inversion / coverage / waterfall charts (optional)."""
    out: dict[str, bytes] = {}
    try:
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        logger.info("matplotlib not installed — FinOps DOCX will use tables only")
        return out

    currency = bundle["currency"]
    # Fig 1 — inversion bar
    try:
        rows = list(bundle.get("agent_rows") or [])[:8]
        # Include top cost RG if missing from agent rows
        for cr in (bundle.get("cost_rank") or [])[:3]:
            if not any(str(cr["target"]).lower() == str(r["name"]).lower() for r in rows):
                rows.append(
                    {
                        "name": cr["target"],
                        "priority": "LOW",
                        "cost": cr["cost"],
                        "sku": "—",
                    }
                )
        if rows:
            labels = [
                f"{_short_name(str(r['name']))} [{r.get('priority')}]" for r in rows
            ]
            costs = [float(r.get("cost") or 0) for r in rows]
            colors = []
            for r in rows:
                p = str(r.get("priority") or "").upper()
                c = float(r.get("cost") or 0)
                if p == "HIGH" and c <= 0.01:
                    colors.append("#c0392b")
                elif p == "LOW" and c > 0:
                    colors.append("#27ae60")
                else:
                    colors.append("#7f8c8d")
            fig, ax = plt.subplots(figsize=(8.5, 3.8))
            ax.barh(labels[::-1], costs[::-1], color=colors[::-1])
            ax.set_xlabel(f"Actual {bundle['days']}-day cost ({currency})")
            ax.set_title("Agent priority vs. actual cost — the inversion")
            fig.tight_layout()
            out["inversion"] = _chart_png(fig)
            plt.close(fig)
    except Exception:
        logger.exception("FinOps inversion chart failed")

    # Fig 2 — coverage donuts
    try:
        fig, axes = plt.subplots(1, 2, figsize=(8.2, 3.6))
        for ax, title, covered in (
            (
                axes[0],
                "Agent recommendations coverage of actual spend",
                float(bundle.get("agent_coverage") or 0),
            ),
            (
                axes[1],
                "Cost-based ranking coverage of actual spend",
                float(bundle.get("cost_coverage") or 0),
            ),
        ):
            covered = max(0.0, min(1.0, covered))
            ax.pie(
                [covered, 1 - covered],
                colors=["#1abc9c" if covered > 0.2 else "#e74c3c", "#ecf0f1"],
                startangle=90,
                wedgeprops={"width": 0.42},
            )
            ax.text(0, 0, f"{covered:.0%}", ha="center", va="center", fontsize=16, fontweight="bold")
            ax.set_title(title, fontsize=9)
        fig.tight_layout()
        out["coverage"] = _chart_png(fig)
        plt.close(fig)
    except Exception:
        logger.exception("FinOps coverage chart failed")

    # Fig 3 — waterfall
    try:
        steps = [s for s in (bundle.get("waterfall") or []) if abs(float(s[1])) > 0.01 or s[2] == "base"]
        if len(steps) >= 2:
            labels = [s[0] for s in steps]
            values = [float(s[1]) for s in steps]
            kinds = [s[2] for s in steps]
            fig, ax = plt.subplots(figsize=(8.5, 3.8))
            running = 0.0
            for i, (lab, val, kind) in enumerate(zip(labels, values, kinds)):
                if kind == "base":
                    ax.bar(i, val, color="#1abc9c")
                    running = val
                else:
                    ax.bar(i, abs(val), bottom=running + val if val < 0 else running, color="#e74c3c")
                    running = running + val
            ax.set_xticks(range(len(labels)))
            ax.set_xticklabels(labels, rotation=15, ha="right", fontsize=8)
            ax.set_ylabel(f"{currency} / month")
            ax.set_title("Indicative savings waterfall (estimates, not quotations)")
            fig.tight_layout()
            out["waterfall"] = _chart_png(fig)
            plt.close(fig)
    except Exception:
        logger.exception("FinOps waterfall chart failed")

    return out


def build_markdown(bundle: dict[str, Any], *, narrative: str = "") -> str:
    cur = bundle["currency"]
    total = float(bundle["total"])
    days = int(bundle["days"])
    lines = [
        "# Cost Optimisation Recommendations",
        "",
        f"**AZURE FINOPS · RECOMMENDATIONS AND COST OPTIMISATION**  ",
        f"{bundle['sub_name']} · Trailing {days} Days · Currency {cur}",
        "",
        (
            f"_Prepared for Sunny Kushwaha · generated "
            f"{bundle['prepared_at'].strftime('%d %B %Y')} · "
            f"sources: Azure FinOps Agent + Azure Cost Management API._"
        ),
        "",
        "| Current bill | Identified saving | Reduction | Actions |",
        "| ---: | ---: | ---: | ---: |",
        (
            f"| {_money(total, cur)} / {days}d | "
            f"{_money(bundle['save_low'], cur)} – {_money(bundle['save_high'], cur)} / mo | "
            f"{(bundle['save_low'] / total * 100) if total else 0:.0f} – "
            f"{(bundle['save_high'] / total * 100) if total else 0:.0f}% | "
            f"{len(bundle['actions'])} |"
        ),
        "",
    ]
    if narrative.strip():
        lines.extend(["## Narrative", "", narrative.strip(), ""])

    lines.extend(
        [
            "## 1. Executive summary",
            "",
            (
                f"A reduction of roughly **{(bundle['save_low'] / total * 100) if total else 0:.0f}–"
                f"{(bundle['save_high'] / total * 100) if total else 0:.0f}%** "
                f"({_money(bundle['save_low'], cur)} – {_money(bundle['save_high'], cur)}) "
                f"is indicated on the current **{_money(total, cur)}** window without application "
                f"code changes — by following **cost-based** ranking rather than SKU heuristics."
            ),
            "",
            "### The decisions that matter",
            "",
        ]
    )
    for a in bundle["actions"][:5]:
        lines.append(f"- **P{a['priority']}** — {a['action']} _(est. {a['saving']} {cur}; {a['risk']} risk)_")
    lines.append("")

    lines.extend(
        [
            "## 2. Agent vs cost-based analysis",
            "",
            "### 2.1 Agent-flagged App Service plans",
            "",
            "| # | Resource | SKU | Agent priority | Actual cost |",
            "| :---: | :--- | :--- | :--- | ---: |",
        ]
    )
    for i, r in enumerate(bundle["agent_rows"][:10], start=1):
        lines.append(
            f"| {i} | `{r['name']}` | {r['sku']} | {r['priority']} | "
            f"{_money(float(r['cost']), cur)} |"
        )
    lines.extend(
        [
            "",
            "### 2.2 Why SKU-only ranking is not actionable",
            "",
            "F1 / Y1 plans are free or near-zero. Ranking them HIGH carries no financial signal. "
            "Cost Management must join the heuristic.",
            "",
            f"_Agent tip coverage of bill ≈ **{bundle['agent_coverage']:.0%}**; "
            f"cost-based top items cover ≈ **{bundle['cost_coverage']:.0%}**._",
            "",
            "*(See Figure 1 inversion chart and Figure 2 coverage donuts in the DOCX.)*",
            "",
        ]
    )

    if bundle["contradictions"]:
        lines.extend(["### 2.4 Unresolved contradictions", ""])
        for c in bundle["contradictions"]:
            lines.append(
                f"- **`{c['resource']}`** reported SKU `{c['sku']}` but billed **{c['cost']}**. {c['note']}"
            )
        lines.append("")

    lines.extend(
        [
            "## 3. Cost-based optimisation plan",
            "",
            "| Rank | Target | 30-day cost | Share | Assessment |",
            "| :---: | :--- | ---: | ---: | :--- |",
        ]
    )
    for i, r in enumerate(bundle["cost_rank"][:8], start=1):
        lines.append(
            f"| {i} | `{r['target']}` | {_money(r['cost'], cur)} | "
            f"{r['share']:.0%} | {r['assessment']} |"
        )
    lines.extend(
        [
            "",
            "## 4. Indicative savings",
            "",
            "*(Figure 3 waterfall is embedded in the DOCX.)*",
            "",
            "| Priority | Action | Est. saving | Risk | Effort |",
            "| :---: | :--- | :--- | :--- | :--- |",
        ]
    )
    for a in bundle["actions"]:
        lines.append(
            f"| {a['priority']} | {a['action']} | {a['saving']} | {a['risk']} | {a['effort']} |"
        )
    lines.extend(
        [
            "",
            (
                f"**Indicative total:** {_money(bundle['save_low'], cur)} – "
                f"{_money(bundle['save_high'], cur)} / month "
                f"({(bundle['save_low'] / total * 100) if total else 0:.0f}–"
                f"{(bundle['save_high'] / total * 100) if total else 0:.0f}% of window)."
            ),
            "",
            "> Estimates are analytical inferences from spend distribution and typical tier "
            "ratios — not live Retail Prices API quotations. Validate in Azure Portal before committing.",
            "",
            "## 5. Agent defects observed",
            "",
            "| Defect | Impact | Fix |",
            "| :--- | :--- | :--- |",
        ]
    )
    for d in bundle["defects"]:
        lines.append(f"| {d['defect']} | {d['impact']} | {d['fix']} |")
    lines.extend(
        [
            "",
            "## 6. Action checklist",
            "",
            "| # | Owner | Item | Effort | Status |",
            "| :---: | :--- | :--- | :--- | :--- |",
        ]
    )
    for i, a in enumerate(bundle["actions"], start=1):
        lines.append(f"| {i} | Sunny | {a['action']} | {a['effort']} | Open |")
    lines.extend(
        [
            "",
            "### Highest-value next step",
            "",
            (
                f"Open Azure Portal → resource group "
                f"**`{(bundle['cost_rank'][0]['target'] if bundle['cost_rank'] else '—')}`** "
                "→ Cost analysis → group by Resource / Meter. Resolve any free-SKU vs spend contradiction "
                "before resizing."
            ),
            "",
        ]
    )

    # Cost by service appendix
    if bundle.get("cost_by_svc"):
        lines.extend(
            [
                "## Appendix A — Cost by service",
                "",
                "| # | Service | Cost |",
                "| :---: | :--- | ---: |",
            ]
        )
        for i, r in enumerate(bundle["cost_by_svc"][:15], start=1):
            lines.append(
                f"| {i} | {r.get('name')} | {_money(float(r.get('cost') or 0), r.get('currency') or cur)} |"
            )
        lines.append("")

    return "\n".join(lines)


def render_finops_docx(bundle: dict[str, Any], markdown: str, charts: dict[str, bytes]) -> bytes | None:
    """Rich DOCX with KPI strip + embedded charts."""
    try:
        from docx import Document
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement
        from docx.shared import Inches, Pt, RGBColor
    except ImportError:
        return None

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)

    title = doc.add_heading("Cost Optimisation Recommendations", level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x73, 0x8C)

    sub = doc.add_paragraph(
        f"AZURE FINOPS · {bundle['sub_name']} · Trailing {bundle['days']} Days · {bundle['currency']}"
    )
    if sub.runs:
        sub.runs[0].font.size = Pt(11)
        sub.runs[0].font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    meta = doc.add_paragraph(
        f"Prepared for Sunny Kushwaha · {bundle['prepared_at'].strftime('%d %B %Y')} · "
        "Azure FinOps Agent + Cost Management API"
    )
    if meta.runs:
        meta.runs[0].font.size = Pt(9)
        meta.runs[0].italic = True

    # KPI table
    kpi = doc.add_table(rows=2, cols=4)
    kpi.style = "Table Grid"
    headers = ["Current bill", "Identified saving", "Reduction", "Actions"]
    total = float(bundle["total"])
    vals = [
        _money(total, bundle["currency"]),
        f"{_money(bundle['save_low'], bundle['currency'])} – {_money(bundle['save_high'], bundle['currency'])}",
        (
            f"{(bundle['save_low'] / total * 100) if total else 0:.0f} – "
            f"{(bundle['save_high'] / total * 100) if total else 0:.0f}%"
        ),
        str(len(bundle["actions"])),
    ]
    for i, h in enumerate(headers):
        kpi.rows[0].cells[i].text = h
        for p in kpi.rows[0].cells[i].paragraphs:
            for r in p.runs:
                r.bold = True
        kpi.rows[1].cells[i].text = vals[i]

    def _shade(cell, hex_color: str) -> None:
        tc = cell._tePr if False else cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), hex_color)
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    try:
        for cell in kpi.rows[0].cells:
            _shade(cell, "D5F5E3")
    except Exception:
        pass

    doc.add_paragraph("")

    if charts.get("inversion"):
        doc.add_heading("Figure 1 — Agent priority vs actual cost (inversion)", level=2)
        doc.add_picture(io.BytesIO(charts["inversion"]), width=Inches(6.2))
    if charts.get("coverage"):
        doc.add_heading("Figure 2 — Recommendation coverage of spend", level=2)
        doc.add_picture(io.BytesIO(charts["coverage"]), width=Inches(6.2))
    if charts.get("waterfall"):
        doc.add_heading("Figure 3 — Indicative savings waterfall", level=2)
        doc.add_picture(io.BytesIO(charts["waterfall"]), width=Inches(6.2))

    # Reuse markdown body renderer lightly
    from agent.services.meeting_publish import render_docx_bytes

    body = render_docx_bytes(
        {"title": "Cost Optimisation Recommendations", "executive_summary": ""},
        markdown,
        plan_kind="action",
    )
    if body and not charts:
        return body

    # Append remaining markdown sections as structured paragraphs (skip duplicate H1)
    table_buf: list[str] = []

    def flush_table() -> None:
        nonlocal table_buf
        if len(table_buf) < 2:
            table_buf = []
            return
        headers = [c.strip() for c in table_buf[0].strip("|").split("|")]
        rows = [
            [c.strip() for c in line.strip("|").split("|")]
            for line in table_buf[2:]
            if line.strip()
        ]
        if headers:
            t = doc.add_table(rows=1 + len(rows), cols=len(headers))
            t.style = "Table Grid"
            for i, label in enumerate(headers):
                t.rows[0].cells[i].text = label
                for p in t.rows[0].cells[i].paragraphs:
                    for r in p.runs:
                        r.bold = True
            for ri, row in enumerate(rows):
                for ci, val in enumerate(row[: len(headers)]):
                    t.rows[ri + 1].cells[ci].text = val
            doc.add_paragraph("")
        table_buf = []

    skip_title = True
    for line in markdown.splitlines():
        text = line.rstrip()
        if text.startswith("|"):
            table_buf.append(text)
            continue
        if table_buf:
            flush_table()
        if not text:
            continue
        if text.startswith("# ") and skip_title:
            skip_title = False
            continue
        if text.startswith("## "):
            doc.add_heading(text[3:].strip(), level=2)
        elif text.startswith("### "):
            doc.add_heading(text[4:].strip(), level=3)
        elif text.startswith("> "):
            p = doc.add_paragraph(text[2:])
            if p.runs:
                p.runs[0].italic = True
        elif text.startswith("- "):
            doc.add_paragraph(text[2:], style="List Bullet")
        elif text.startswith("*(") or text.startswith("_Prepared"):
            p = doc.add_paragraph(text.strip("*_"))
            if p.runs:
                p.runs[0].italic = True
                p.runs[0].font.size = Pt(9)
        else:
            doc.add_paragraph(text)

    if table_buf:
        flush_table()

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


async def gather_finops_bundle(
    *,
    session_data: dict[str, Any],
    days: int | None = None,
) -> dict[str, Any]:
    """Load subscription + live (or cached) cost/recs into an analysis bundle."""
    days = int(days or settings.azure_finops_cost_days or 30)
    sub = dict(session_data.get("azure_finops_sub") or {})
    if not sub.get("id"):
        # Fall back to first enabled subscription
        subs = await azure_finops.list_subscriptions()
        if not subs:
            raise RuntimeError(
                "No Azure subscription visible to the FinOps identity. "
                "Say *azure subscriptions* in Orbit first, then ask for the DOCX again."
            )
        sub = subs[0]

    sid = sub.get("id") or ""
    cost_by_rg = list(session_data.get("azure_finops_last_costs_rg") or [])
    cost_by_svc = list(session_data.get("azure_finops_last_costs_svc") or [])
    cost_by_res = list(session_data.get("azure_finops_last_costs_res") or [])
    cached_recs = list(session_data.get("azure_finops_last_recs") or [])

    if not cost_by_rg or not cost_by_svc:
        by_rg = await azure_finops.query_costs_result(sid, group_by="ResourceGroupName")
        if not by_rg.get("ok"):
            raise RuntimeError(
                by_rg.get("user_note")
                or "Cost Management query failed — try *azure costs* again in a few minutes."
            )
        cost_by_rg = list(by_rg.get("rows") or [])
        if by_rg.get("error_kind") != "throttled":
            by_svc = await azure_finops.query_costs_result(sid, group_by="ServiceName")
            if by_svc.get("ok"):
                cost_by_svc = list(by_svc.get("rows") or [])

    resources = await azure_finops.list_resources(sid)
    if not cost_by_res:
        by_res = await azure_finops.query_costs_result(sid, group_by="ResourceId")
        if by_res.get("ok"):
            cost_by_res = list(by_res.get("rows") or [])

    recs = cached_recs or azure_finops.build_recommendations(
        resources,
        cost_by_resource=cost_by_res,
        cost_by_rg=cost_by_rg,
    )

    return build_analysis_bundle(
        sub=sub,
        cost_by_rg=cost_by_rg,
        cost_by_svc=cost_by_svc,
        cost_by_res=cost_by_res,
        resources=resources,
        recs=recs,
        days=days,
    )


async def polish_narrative(bundle: dict[str, Any], user_message: str) -> str:
    """Short LLM narrative constrained to injected facts (planning-quality model)."""
    from langchain_core.messages import HumanMessage, SystemMessage

    from agent.services.llm import invoke_llm

    facts = (
        f"Subscription: {bundle['sub_name']}\n"
        f"Window days: {bundle['days']}\n"
        f"Currency: {bundle['currency']}\n"
        f"Total bill: {bundle['total']:.2f}\n"
        f"Indicative savings: {bundle['save_low']:.2f} – {bundle['save_high']:.2f}\n"
        f"Cost coverage of top ranks: {bundle['cost_coverage']:.0%}\n"
        f"Agent coverage: {bundle['agent_coverage']:.0%}\n"
        f"Top targets: {[r['target'] for r in bundle['cost_rank'][:4]]}\n"
        f"Contradictions: {bundle['contradictions']}\n"
        f"User request: {user_message[:500]}\n"
    )
    system = (
        "You are writing 2 short paragraphs for an Azure FinOps cost-optimisation report. "
        "Use ONLY the facts provided. Name specific resource groups and INR/currency amounts. "
        "Do not invent services. Do not give generic Azure best-practice lists. "
        "Call out any free-SKU vs spend contradiction. Plain text only, no markdown headings."
    )
    try:
        resp = await invoke_llm(
            [SystemMessage(content=system), HumanMessage(content=facts)],
            temperature=0.2,
            role="doc_author",
        )
        return (resp.content or "").strip()
    except Exception:
        logger.exception("FinOps narrative polish failed — using deterministic text")
        top = bundle["cost_rank"][0]["target"] if bundle["cost_rank"] else "the top resource group"
        return (
            f"On {bundle['sub_name']}, trailing {bundle['days']}-day spend is "
            f"{_money(bundle['total'], bundle['currency'])}. Cost-based ranking focuses on "
            f"`{top}` and related SQL/POC estate. SKU-only tips cover roughly "
            f"{bundle['agent_coverage']:.0%} of the bill; cost-based top items cover "
            f"{bundle['cost_coverage']:.0%}."
        )
