"""Data Analyst Agent — Excel → executive analytics workbook.

Feature: ENABLE_DATA_ANALYST_AGENT default false. Isolated from Dev/WhatsApp coding
and from Doc RAG ingest. Uses Doc Library catalog picks when present.
"""

from __future__ import annotations

import csv
import io
import json
import re
from datetime import UTC, datetime
from pathlib import Path

from langchain_core.messages import HumanMessage, SystemMessage
from openpyxl import Workbook, load_workbook

from agent.agents.state import AgentState
from agent.config import settings
from agent.core.agent_availability import data_analyst_offline
from agent.core.logging import get_logger
from agent.core.session import get_session
from agent.services import graph_docs, ms_graph
from agent.services.excel_workbook import (
    build_analytics_workbook,
    parse_insights_json,
    profile_workbook,
)
from agent.services.llm import invoke_llm

logger = get_logger(__name__)

AGENT_NAME = "data_analyst_agent"
PROMPT_PATH = Path(__file__).parent.parent / "prompts" / "excel_analyst_system.txt"
_NUMS_RE = re.compile(r"\b(\d+)\b")
_MAX_BYTES = 12 * 1024 * 1024

_DISABLED = data_analyst_offline()


def _is_excel_row(row: dict) -> bool:
    ext = (row.get("extension") or "").lower().lstrip(".")
    title = (row.get("title") or "").lower()
    return ext in ("xlsx", "xlsm", "csv") or title.endswith((".xlsx", ".xlsm", ".csv"))


async def analyze_excel(state: AgentState) -> AgentState:
    if not settings.enable_data_analyst_agent:
        return {
            **state,
            "status": "skipped",
            "handled_by": AGENT_NAME,
            "notification_text": _DISABLED,
        }

    if not ms_graph.graph_configured():
        return {
            **state,
            "status": "failed",
            "handled_by": AGENT_NAME,
            "notification_text": (
                "*Data Analyst Agent* — Graph credentials missing.\n"
                "See `docs/DATA_ANALYST_AGENT.md`."
            ),
        }

    msg = state.get("user_message") or ""
    phone = state.get("whatsapp_phone") or ""
    catalog: list = []
    if phone:
        try:
            catalog = list(((await get_session(phone)).get("data") or {}).get("doc_catalog") or [])
        except Exception:
            catalog = []
    if not catalog:
        try:
            catalog = await graph_docs.list_knowledge_catalog()
        except Exception as exc:
            return {
                **state,
                "status": "failed",
                "handled_by": AGENT_NAME,
                "notification_text": (
                    f"*Data Analyst Agent* — no catalog and list failed: {exc}\n"
                    "Say *list my documents*, then *convert excel &lt;n&gt;*."
                ),
                "error": str(exc),
            }

    picks = {int(n) for n in _NUMS_RE.findall(msg)}
    excel_rows = [r for r in catalog if _is_excel_row(r)]
    if picks:
        selected = [r for r in catalog if int(r.get("pick") or 0) in picks]
    elif len(excel_rows) == 1:
        selected = excel_rows
    else:
        selected = []

    if not selected:
        hint = ", ".join(str(r.get("pick")) for r in excel_rows[:8]) or "none on this list"
        return {
            **state,
            "status": "failed",
            "handled_by": AGENT_NAME,
            "notification_text": (
                "*Data Analyst Agent* — pick an Excel/CSV file from the library.\n"
                f"Try *convert excel 6*. Excel picks on the last list: {hint}."
            ),
        }

    entry = selected[0]
    if not _is_excel_row(entry):
        return {
            **state,
            "status": "failed",
            "handled_by": AGENT_NAME,
            "notification_text": (
                f"*{entry.get('title')}* is not Excel/CSV. "
                "Pick an `.xlsx` / `.xlsm` / `.csv` row."
            ),
        }

    raw, _ct = await graph_docs.fetch_document_bytes(entry)
    if not raw:
        return {
            **state,
            "status": "failed",
            "handled_by": AGENT_NAME,
            "notification_text": (
                f"*Data Analyst Agent* could not download *{entry.get('title')}*."
            ),
        }
    if len(raw) > _MAX_BYTES:
        return {
            **state,
            "status": "failed",
            "handled_by": AGENT_NAME,
            "notification_text": (
                f"*{entry.get('title')}* is larger than 12 MB. "
                "Split the workbook or filter rows, then retry."
            ),
        }

    title = entry.get("title") or "workbook.xlsx"
    insights: dict = {}
    try:
        if title.lower().endswith(".csv"):
            tmp = Workbook()
            ws = tmp.active
            ws.title = "RawCSV"
            text = raw.decode("utf-8-sig", errors="replace")
            for r_i, row in enumerate(csv.reader(io.StringIO(text)), start=1):
                for c_i, val in enumerate(row, start=1):
                    ws.cell(r_i, c_i, val)
            profile = profile_workbook(tmp)
        else:
            profile = profile_workbook(load_workbook(io.BytesIO(raw), data_only=True))
        system = PROMPT_PATH.read_text(encoding="utf-8")
        resp = await invoke_llm(
            [
                SystemMessage(content=system),
                HumanMessage(
                    content="Dataset profile JSON:\n" + json.dumps(profile, default=str)[:12000]
                ),
            ],
            temperature=0.1,
            role="planning",
        )
        insights = parse_insights_json(getattr(resp, "content", "") or "")
    except Exception:
        logger.exception("Data Analyst insight model failed — continuing with heuristics")

    out_bytes = build_analytics_workbook(raw, title=title, insights=insights, filename=title)
    stamp = datetime.now(UTC).strftime("%Y%m%d-%H%M%S")
    stem = re.sub(r"[^\w.\-]+", "-", title.rsplit(".", 1)[0])[:40].strip("-") or "workbook"
    out_name = f"{stem}-analytics-{stamp}.xlsx"
    url = await graph_docs.upload_site_file(
        folder="Analytics",
        filename=out_name,
        content=out_bytes,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    if not url:
        return {
            **state,
            "status": "failed",
            "handled_by": AGENT_NAME,
            "notification_text": (
                "*Data Analyst Agent* built the workbook but SharePoint upload failed.\n"
                "Check Graph site write permissions (same as Docs Agent)."
            ),
        }

    n_insights = len(insights.get("insights") or [])
    note = (
        "*Data Analyst Agent* — executive analytics workbook published\n\n"
        f"Source: *{title}*\n"
        f"Output: [{out_name}]({url})\n"
        f"Folder: SharePoint **Analytics**\n"
        f"Insights generated: {n_insights} (grounded in the profile; original sheets copied to Raw_*)\n\n"
        "Open the file for README → Dashboard → Insights → Data Quality → Interactive Analysis.\n"
        "Original library file was not modified.\n\n"
        "**Next:** *list my documents* · *convert excel &lt;n&gt;* for another sheet"
    )
    return {
        **state,
        "status": "completed",
        "handled_by": AGENT_NAME,
        "notification_text": note,
    }
