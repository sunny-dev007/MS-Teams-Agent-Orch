"""Data Analyst workbook builder — no Graph/LLM."""

from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook

from agent.services.excel_workbook import build_analytics_workbook, parse_insights_json


def _sample_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Revenue", "Month"])
    ws.append(["East", 100, "2026-01"])
    ws.append(["West", 80, "2026-01"])
    ws.append(["East", 120, "2026-02"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_build_analytics_workbook_layers():
    raw = _sample_xlsx()
    insights = {
        "business_summary": "Regional sales sample.",
        "insights": [
            {
                "key_insight": "East leads revenue in the sample.",
                "driver": "Two East rows vs one West row.",
                "business_implication": "Validate with a fuller extract before investment.",
            }
        ],
        "quality_flags": [],
        "sensitive_fields": [],
    }
    out = build_analytics_workbook(raw, title="Sales.xlsx", insights=insights, filename="Sales.xlsx")
    wb = load_workbook(io.BytesIO(out))
    names = set(wb.sheetnames)
    assert "README" in names
    assert "Executive_Dashboard" in names
    assert "Insights" in names
    assert "Data_Quality" in names
    assert "Interactive_Analysis" in names
    assert "Configuration" in names
    assert any(n.startswith("Raw_") for n in names)
    raw_sheet = next(n for n in wb.sheetnames if n.startswith("Raw_"))
    assert wb[raw_sheet]["A1"].value == "Region"
    assert wb[raw_sheet]["A2"].value == "East"


def test_parse_insights_json_fenced():
    data = parse_insights_json('```json\n{"business_summary": "ok", "insights": []}\n```')
    assert data.get("business_summary") == "ok"
